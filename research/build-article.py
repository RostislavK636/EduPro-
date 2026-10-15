"""Run with the bundled workspace Python. Input evidence must already exist."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent.parent
results=json.loads((root/'research/results/technical.json').read_text(encoding='utf-8'))
sim=json.loads((root/'research/results/simulation.json').read_text(encoding='utf-8'))
out=root.parent/'outputs';out.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman');st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme']:
        st._element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1.25)
for style in doc.styles:
    for border in list(style._element.xpath('.//w:pBdr')):border.getparent().remove(border)
content=[]
def p(text,bold=False,center=False,style=None):
    content.append(text);para=doc.add_paragraph(style=style);para.add_run(text).bold=bold
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if center or style=='Title':para.paragraph_format.first_line_indent=Cm(0)
    return para
def h(text):
    para=p(text,True);para.paragraph_format.keep_with_next=True
p('УДК 004.8:37')
p('EduPilot как образовательная платформа с объяснимым адаптивным тренажёром',True,True,'Title')
p('Автор 1 ____________________  IT-специалист',center=True)
p('Автор 2 ____________________',center=True)
p('Автор 3 ____________________',center=True)
p('Организация ____________________  Москва',center=True)
p('Научный руководитель ____________________',center=True)
p('Аннотация. Представлено расширение образовательной платформы EduPilot детерминированным адаптивным тренажёром и аналитикой учебных попыток. Описаны диагностика, выбор сложности и журнал оснований каждого решения. Корректность проверяется техническими сценариями и воспроизводимой симуляцией двух режимов на общем банке заданий. Педагогическая эффективность и качество реальной генеративной модели пока не измерены.')
p('Ключевые слова: адаптивное обучение, образовательная платформа, искусственный интеллект, учебная аналитика, объяснимость, программное тестирование.')
h('Введение')
p('Образовательной платформе недостаточно генерировать текст: учащемуся нужна последовательность упражнений с обратной связью, а преподавателю — наблюдаемые основания поддержки. Модели knowledge tracing используют историю взаимодействия для представления прогресса [1, 2]. Исследования репетиторских систем показывают необходимость различать механизмы помощи и условия сравнения [3, 4]. Результаты специально спроектированного AI-тьютора [5] нельзя переносить на любой чат-бот. Исследование школьной математики [6] дополнительно мотивирует самостоятельное оценивание после практики с помощником. Возможности и ограничения больших языковых моделей рассматриваются в [7].')
p('Цель работы — дополнить существующую платформу объяснимым адаптивным циклом, сохранив её материалы, классы и интерфейс интерактивного созвездия. Объект — программная поддержка учебной практики; предмет — поведение правил подбора сложности по ответам. Техническая гипотеза состоит в воспроизводимости переходов и сохранении их оснований. Гипотеза о лучшем самостоятельном освоении требует отдельной педагогической апробации.')
h('Практическая новизна и методы')
p('Практический вклад заключается в объединении снимка вопроса, идемпотентной записи ответа и доступной преподавателю истории решений внутри действующей платформы. Новизна универсального алгоритма не заявляется. Использованы аудит исходного кода, тесты чистых функций, интеграционные запросы к изолированной SQLite и сценарии реального браузера. Алгоритм сравнивается с фиксированным выбором на одном банке при одинаковых правилах проверки и критерии освоения.')
h('Архитектура и алгоритм')
p('Клиент Next/React взаимодействует с сервером через авторизованный API; Prisma сохраняет данные в SQLite либо D1. Добавочная миграция не заменяет прежние таблицы. Учебная сессия фиксирует тему, режим и конфигурацию. Для каждого шага сохраняется неизменяемый снимок задания. Уникальность попытки по шагу предотвращает двойной ответ: повтор того же запроса возвращает существующий результат, конфликтующий ответ отклоняется. Текущий эталон не передаётся ученику до сдачи. Учитель видит только учащихся своих классов.')
p('Банк содержит 72 задания по четырём темам математики и информатики, три уровня сложности. Диагностика включает пять вопросов уровней 1, 1, 2, 2, 3. Менее 50% верных ответов задают базовый уровень, 50–79% — средний, от 80% — продвинутый. В десяти заданиях практики три верных подряд повышают уровень, две ошибки снижают; границы равны 1 и 3. Серия сбрасывается после срабатывания. Коэффициенты и пороги являются настройками проекта, а не научно установленным оптимумом.')
p('Выбор учитывает расстояние до целевого уровня, предыдущие предъявления и метки повторных ошибок. Последние три задания исключаются при наличии альтернатив. Mastery — сглаженная взвешенная доля с начальными псевдосчётчиками успеха и ошибки по единице. Вес ответа равен уровню, делённому на номер повторной попытки того же вопроса. Более сложные решения дают больший вклад, повторные — меньший. Критерий «освоено» требует не менее шести ответов, mastery не ниже 0,8 и двух верных ответов уровней 2–3. Показатель эвристический и не является валидированной оценкой знания.')
p('Правила не обучаются на данных и не называются машинным обучением. Генеративный AI опционален: он объясняет тему, предлагает разбор ошибки и создаёт черновик. Ключ хранится на сервере, профиль пользователя в запрос не входит. Учебный свободный текст требует предварительного контроля персональных сведений. Задание проходит состояния draft, pending, approved либо rejected; правка снимает одобрение. Проверка JSON не заменяет предметную экспертизу учителя. При отсутствии ключа демо-режим обозначается явно; локальная проверка ответа продолжает работать.')
h('Технические эксперименты')
p(f"На запуске {results['generatedAt'][:10]} выполнено {results['total']} автоматических проверок, пройдено {results['passed']}, ошибок {results['failed']}. Среди них {results['predefinedLearningScenarios']} заранее определённых сценария ядра и API: границы диагностики, серии ответов, повторы, смена темы, восстановление, параллельная сдача, межпользовательский доступ и некорректный ответ провайдера. Исходная платформа до расширения проходила 14 тестов. Эти числа характеризуют программные проверки, а не количество участников исследования. Полные результаты, версии и диагностика сохраняются автоматически в JSON/CSV.")
p('Симуляция включает три заданных профиля ответа, четыре темы и два режима: 24 траектории по 15 ответов, всего 360. Реальных учащихся нет. Seed и функция ответа фиксированы; вероятность зависит от сложности и экспозиции, но не от названия режима. Первые пять заданий и ответы одинаковы в каждой паре. Далее фиксированный режим циклически предъявляет сложности, адаптивный следует текущему уровню. В таблице приведены фактически рассчитанные агрегаты по четырём темам каждого профиля.')
table=doc.add_table(rows=1, cols=4);table.autofit=False
headers=['Профиль','Режим','Верно / всего','Смены уровня*']
for c,text in zip(table.rows[0].cells,headers):c.text=text
names={'stable-capable':'Устойчивый','recurring-errors':'Повторные ошибки','improving':'Рост вероятности'}
for row in sim['comparisons']:
    values=[names[row['profileId']],'Фикс.' if row['mode']=='fixed' else 'Адапт.',f"{row['accuracy']['numerator']} / {row['accuracy']['denominator']}",str(row['practiceLevelChanges'])]
    for c,text in zip(table.add_row().cells,values):c.text=text
for row in table.rows:
    for i,c in enumerate(row.cells):
        c.width=Cm([5.2,3,3.7,4.1][i]);tcPr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
        tcPr.append(borders)
        for para in c.paragraphs:para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
for c in table.rows[0].cells:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EBEF');c._tc.get_or_add_tcPr().append(shade)
    for run in c.paragraphs[0].runs:run.bold=True
p('* Только изменения в практике, без назначения уровня по диагностике.')
p('Результаты не дают универсального преимущества адаптации. Устойчивый профиль получает больше сложных повторов, а профиль ошибок — более доступные задания. Разная экспозиция влияет на долю верных ответов; она не измеряет прирост знаний. Количество заданий до эвристического освоения и полные трассы доступны в машинном отчёте. Вероятности профиля заданы исследователем, поэтому симуляция проверяет поведение механизма, а не педагогический эффект.')
h('Выводы и ограничения')
p('Реализован автономный цикл диагностики, практики и обратной связи с сохранением оснований адаптации. Аналитика строится по принятым попыткам, показывает знаменатели и предупреждает о малой выборке; demo хранится отдельно. Ограничения — небольшой неоткалиброванный банк, эвристические параметры, отсутствие реальных pre/post-данных и непроверенное качество Live AI. Подготовлен протокол добровольной апробации с самостоятельными сопоставимыми pre/post-тестами, учётом времени, повторных ошибок и удобства. До её проведения образовательная эффективность не заявляется.')
h('Список литературы')
refs=[
'Corbett A. T., Anderson J. R. Knowledge tracing: Modeling the acquisition of procedural knowledge // User Modeling and User-Adapted Interaction. 1994. Vol. 4. P. 253–278. DOI: 10.1007/BF01099821.',
'Deep Knowledge Tracing / C. Piech, J. Bassen, J. Huang [et al.] // Advances in Neural Information Processing Systems 28. 2015. URL: https://proceedings.neurips.cc/paper/2015/hash/bac9162b47c56fc8a4d2a519803d51b3-Abstract.html (дата обращения: 15.10.2026).',
'VanLehn K. The Relative Effectiveness of Human Tutoring, Intelligent Tutoring Systems, and Other Tutoring Systems // Educational Psychologist. 2011. Vol. 46, no. 4. P. 197–221. DOI: 10.1080/00461520.2011.611369.',
'Kulik J. A., Fletcher J. D. Effectiveness of Intelligent Tutoring Systems: A Meta-Analytic Review // Review of Educational Research. 2016. Vol. 86, no. 1. P. 42–78. DOI: 10.3102/0034654315581420.',
'AI tutoring outperforms in-class active learning: an RCT introducing a novel research-based design in an authentic educational setting / G. Kestin, K. Miller, A. Klales [et al.] // Scientific Reports. 2025. Vol. 15. Art. 17458. DOI: 10.1038/s41598-025-97652-6.',
'Generative AI without guardrails can harm learning: Evidence from high school mathematics / H. Bastani, O. Bastani, A. Sungu [et al.] // PNAS. 2025. Vol. 122, no. 26. Art. e2422633122. DOI: 10.1073/pnas.2422633122.',
'ChatGPT for good? On opportunities and challenges of large language models for education / E. Kasneci, K. Sessler, S. Küchemann [et al.] // Learning and Individual Differences. 2023. Vol. 103. Art. 102274. DOI: 10.1016/j.lindif.2023.102274.'
]
for n,r in enumerate(refs,1):
    para=p(f'{n}. {r}');para.paragraph_format.first_line_indent=Cm(0)
doc.core_properties.title='EduPilot как образовательная платформа с объяснимым адаптивным тренажёром';doc.core_properties.author=''
for para in list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for c in row.cells for p in c.paragraphs]:
    for border in list(para._p.xpath('.//w:pBdr')):border.getparent().remove(border)
    for run in para.runs:
        run.font.name='Times New Roman';run.font.size=Pt(14);run.font.color.rgb=RGBColor(0,0,0)
        fonts=run._element.get_or_add_rPr().rFonts
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme']:fonts.attrib.pop(qn('w:'+attr),None)
doc.save(out/'EduPilot-article.docx')
(root/'research/ARTICLE.md').write_text('\n\n'.join(content),encoding='utf-8')
print(json.dumps({'output':str(out/'EduPilot-article.docx'),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'references':len(refs),'recentReferences':3,'pageLimit':5,'renderRequired':True},ensure_ascii=False))
